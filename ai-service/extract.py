"""Extract plain text from uploaded files (in memory, nothing is saved to disk).

Arabic PDFs are lossy: PDF stores text in visual order, so numbers can drift
within a sentence. DOCX / TXT / MD keep the logical order and are preferred.
"""
import io
import unicodedata

from docx import Document
from pypdf import PdfReader

MAX_CHARS = 300_000


def _from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _from_pdf(data)
    elif name.endswith(".docx"):
        text = _from_docx(data)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8-sig", errors="replace")
    else:
        raise ValueError("Unsupported file type. Use DOCX, TXT, MD or PDF.")
    # NFKC turns Arabic presentation forms (connected letter shapes) into normal letters
    text = unicodedata.normalize("NFKC", text)
    return text.strip()[:MAX_CHARS]